from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from pypdf import PdfReader


ROOT = Path(r"C:\Suivi-reclamation")
SRC = Path(r"D:\Mes données\ANBG stage\Rapport_stage_2026_FANEL.docx")
TEMPLATE = Path(r"D:\RAPPORT DE STAGE_Gabarit_mise en page_Word_IITG_2026 (1).pdf")
OUT = ROOT / "output" / "documents" / "Rapport_stage_ANBG_IITG_adapte.docx"
ASSETS = ROOT / "tmp" / "rapport_stage"
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document(SRC)
original = list(doc.paragraphs)
anchor = original[0]


def tune_style(name: str, size: float, *, bold: bool = False, before: int = 0,
               after: int = 6, keep_next: bool = False) -> None:
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = None
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.5 if name == "Normal" else 1.2
    fmt.keep_with_next = keep_next


tune_style("Normal", 12, after=6)
tune_style("Title", 18, bold=True, before=12, after=14, keep_next=True)
tune_style("Heading 1", 15, bold=True, before=18, after=12, keep_next=True)
tune_style("Heading 2", 14, bold=True, before=16, after=10, keep_next=True)
tune_style("Heading 3", 13, bold=True, before=14, after=8, keep_next=True)
tune_style("Heading 4", 12, bold=True, before=10, after=6, keep_next=True)
tune_style("Heading 5", 11, bold=True, before=8, after=5, keep_next=True)

for st_name in ("Report Front", "Report Chapter Label", "Report Code"):
    if st_name not in doc.styles:
        doc.styles.add_style(st_name, WD_STYLE_TYPE.PARAGRAPH)

front_style = doc.styles["Report Front"]
front_style.base_style = doc.styles["Normal"]
front_style.font.name = "Arial"
front_style.font.size = Pt(12)
front_style.paragraph_format.space_after = Pt(9)
front_style.paragraph_format.line_spacing = 1.25

code_style = doc.styles["Report Code"]
code_style.font.name = "Consolas"
code_style.font.size = Pt(8.5)
code_style.paragraph_format.space_after = Pt(5)
code_style.paragraph_format.line_spacing = 1.0

for section in doc.sections:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(2.45)
    section.bottom_margin = Cm(2.45)
    section.left_margin = Cm(2.55)
    section.right_margin = Cm(2.45)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = True


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " " + instruction + " "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


header = doc.sections[0].header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.style = doc.styles["Normal"]
header.add_run("Rapport de stage — ANBG  |  ")
add_field(header, "PAGE", "1")
for run in header.runs:
    run.font.size = Pt(9)


def front(text: str = "", style: str = "Report Front", *, center: bool = False,
          page_break: bool = False, bold: bool = False, after: int | None = None):
    p = anchor.insert_paragraph_before(style=style)
    p.paragraph_format.page_break_before = page_break
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        r = p.add_run(text)
        if bold:
            r.bold = True
    return p


def front_heading(text: str, *, page_break: bool = True):
    p = front(text, "Title", center=True, page_break=page_break)
    return p


logo_path = ASSETS / "iitg_logo_reference.jpg"
try:
    logo_path.write_bytes(PdfReader(str(TEMPLATE)).pages[0].images[0].data)
except Exception:
    logo_path = None

# Preamble: the template is used for typography and page rhythm, while the
# methodology controls which components belong in the report.
front("INSTITUT INTERNATIONAL DE TECHNOLOGIE ET DE GESTION (IITG)", center=True, bold=True, after=20)
if logo_path and logo_path.exists():
    p = front(center=True, after=28)
    p.add_run().add_picture(str(logo_path), width=Cm(3.1))
front("RAPPORT DE STAGE", "Title", center=True, after=18)
front("Conception et réalisation d’une application de gestion des réclamations\nau sein de l’Agence nationale des bourses du Gabon (ANBG)", center=True, bold=True, after=28)
front("Présenté par : [Nom et prénom à renseigner]", center=True)
front("Formation / spécialité : [à renseigner]", center=True)
front("Structure d’accueil : Agence nationale des bourses du Gabon (ANBG)", center=True)
front("Encadreur académique : [à renseigner]", center=True)
front("Encadreur professionnel : [à renseigner]", center=True)
front("Période de stage : [à renseigner]", center=True, after=25)
front("Année académique : 2025–2026", center=True)

# Blank verso, followed by the formal title page required by the methodology.
front("", page_break=True)
front("INSTITUT INTERNATIONAL DE TECHNOLOGIE ET DE GESTION (IITG)", center=True, page_break=True, bold=True, after=16)
if logo_path and logo_path.exists():
    p = front(center=True, after=20)
    p.add_run().add_picture(str(logo_path), width=Cm(2.8))
front("RAPPORT DE STAGE", "Title", center=True)
front("Conception et réalisation d’une application de gestion des réclamations à l’ANBG", center=True, bold=True, after=28)
front("Rapport présenté en vue de l’obtention du diplôme : [à renseigner]", center=True)
front("Par : [Nom et prénom à renseigner]", center=True)
front("Lieu et période de stage : ANBG — [à renseigner]", center=True, after=22)
front("Encadreur académique : [à renseigner]", center=True)
front("Encadreur professionnel : [à renseigner]", center=True)
front("Année académique 2025–2026", center=True, after=0)

front_heading("DÉDICACE")
front("[À personnaliser par l’auteur du rapport.]")

front_heading("REMERCIEMENTS")
front("[À personnaliser : responsables de l’ANBG, encadreurs et personnes ayant accompagné le stage.]")

front_heading("AVANT-PROPOS")
front("Le présent rapport s’inscrit dans le cadre du stage prévu par la formation [intitulé à renseigner] de l’Institut international de technologie et de gestion. Effectué au sein de l’Agence nationale des bourses du Gabon, il rend compte de l’environnement professionnel découvert, des tâches réalisées et du projet de gestion des réclamations étudié pendant cette période.")

front_heading("LISTE DES SIGLES ET ABRÉVIATIONS")
abbreviations = [
    ("ANBG", "Agence nationale des bourses du Gabon"),
    ("DSIC", "Direction des systèmes d’information et de la communication"),
    ("SCIQ", "Service du contrôle interne et de la qualité"),
    ("UCAS", "Unité courrier, accueil et sécurité"),
    ("UML", "Unified Modeling Language (langage de modélisation unifié)"),
    ("MCD", "Modèle conceptuel de données"),
    ("MVC", "Modèle–Vue–Contrôleur"),
    ("SLA", "Service Level Agreement (délai de traitement cible)"),
]
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.rows[0].cells[0].text = "Sigle"
table.rows[0].cells[1].text = "Signification"
for abbr, meaning in abbreviations:
    cells = table.add_row().cells
    cells[0].text = abbr
    cells[1].text = meaning
anchor._p.addprevious(table._tbl)

front_heading("SOMMAIRE")
sommaire = front()
add_field(sommaire, 'TOC \\o "1-3" \\h \\z \\u', "Mettre à jour le sommaire dans Word.")

def set_heading(paragraph, text: str, level: int, *, page_break: bool = False):
    paragraph.text = text
    paragraph.style = doc.styles[f"Heading {level}"]
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.keep_with_next = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.bold = True
        run.font.size = Pt({1: 15, 2: 14, 3: 13, 4: 12, 5: 11}[level])


# Make chapter and section hierarchy faithful to the two-part methodology.
set_heading(original[0], "INTRODUCTION GÉNÉRALE", 1, page_break=True)
part1 = original[30].insert_paragraph_before("PREMIÈRE PARTIE : PRÉSENTATION DE L’ANBG ET DU STAGE")
set_heading(part1, part1.text, 1, page_break=True)
set_heading(original[30], "CHAPITRE I : PRÉSENTATION DE LA STRUCTURE D’ACCUEIL", 2)
original[31].text = ""
set_heading(original[35], "Section 1 : Présentation générale de l’ANBG", 3)
set_heading(original[56], "Section 2 : Organisation et fonctionnement de l’ANBG", 3)
set_heading(original[83], "Conclusion du chapitre I", 4)

set_heading(original[89], "CHAPITRE II : PRÉSENTATION DU SERVICE D’ACCUEIL ET DES TÂCHES EFFECTUÉES", 2, page_break=True)
original[90].text = ""
set_heading(original[93], "Section 1 : Présentation du service d’accueil", 3)
set_heading(original[116], "Section 2 : Tâches effectuées", 3)
set_heading(original[123], "Conclusion du chapitre II", 4)

part2 = original[129].insert_paragraph_before("DEUXIÈME PARTIE : ÉTUDE ET RÉALISATION DE L’APPLICATION")
set_heading(part2, part2.text, 1, page_break=True)
set_heading(original[129], "CHAPITRE III : CADRE D’ÉTUDE", 2)
original[130].text = ""
set_heading(original[134], "Section 1 : Techniques d’investigation", 3)
set_heading(original[157], "Section 2 : Présentation du problème et modélisation du système", 3)
set_heading(original[247], "3.8.1 Description des cas d’utilisation", 4)

set_heading(original[299], "CHAPITRE IV : RÉSULTATS, ANALYSE ET SUGGESTIONS", 2, page_break=True)
original[300].text = ""
set_heading(original[303], "Section 1 : Présentation et mise en œuvre de la solution", 3)
set_heading(original[359], "4.2.3 Framework Laravel", 5)
set_heading(original[409], "4.5.2.6 Sécurisation du serveur", 5)

# Clarify roles in the pre-existing Gmail workflow without changing its facts.
for p in original[159:183]:
    if "constituent la cellule" in p.text.lower():
        p.text = ("L’UCAS, unité chargée du courrier, de l’accueil et de la sécurité, assure la réception "
                  "des messages envoyés à la cellule réclamations. Le SCIQ accède également à cette "
                  "boîte pour le contrôle et le reporting. Les chefs de service et, le cas échéant, "
                  "leurs agents interviennent par transfert des courriels, sans accès direct à la boîte.")

# Normalize original hierarchy and body typography, preserving figures and tables.
for p in original:
    if p in (original[31], original[90], original[130]):
        continue
    txt = p.text.strip()
    if p.style.name.startswith("Heading"):
        continue
    if not txt:
        continue
    if "\n" in txt and ("sudo " in txt or "php8.3" in txt or "server {" in txt):
        p.style = doc.styles["Report Code"]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        continue
    if re.match(r"^\d+\.\d+(?:\.\d+)*\s*\.?\s+[A-Za-zÀ-ÿ]", txt) and len(txt) < 145:
        depth = txt.split(" ", 1)[0].count(".")
        level = 4 if depth <= 2 else 5
        set_heading(p, txt, level)
        continue
    if p.style.name == "Normal":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        if run.text and not run.font.name == "Consolas":
            run.font.name = "Arial"
            if p.style.name == "Normal":
                run.font.size = Pt(12)

for table in doc.tables:
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.5)
                    if i == 0:
                        run.bold = True


def add_paragraph(text: str = "", style: str = "Normal", *, align=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    if align is not None:
        p.alignment = align
    elif style == "Normal":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_heading(text: str, level: int, *, page_break: bool = False):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    p.paragraph_format.page_break_before = page_break
    return p


add_heading("Section 2 : Analyse des résultats et suggestions", 3)
add_heading("4.6 Analyse des résultats", 4)
add_paragraph("La solution réunit dans une même application les étapes auparavant dispersées dans les échanges Gmail et le fichier Excel : dépôt de la réclamation, qualification par l’UCAS, affectation au service compétent, suivi des interventions et consultation des données de pilotage. Les accès sont différenciés selon les rôles. Cette centralisation rend le parcours d’un dossier plus lisible et facilite la recherche de son historique.")
add_paragraph("Les diagrammes de cas d’utilisation, de classes et le modèle conceptuel de données présentés au chapitre III ont servi à relier les besoins recueillis aux fonctions mises en œuvre. Les écrans et modules décrits dans la première section de ce chapitre montrent notamment le suivi des demandes, la gestion des affectations, les notifications et l’export des informations utiles au contrôle interne. Le déploiement décrit dans le rapport établit la faisabilité technique de la solution ; il ne suffit toutefois pas, à lui seul, à mesurer son effet sur les délais réels de traitement.")
add_paragraph("Une vérification technique complémentaire de la version disponible de l’application a porté sur trois groupes de tests automatisés : le parcours des réclamations et les courriels associés, le calcul du délai de traitement en heures ouvrées, ainsi que l’export Word de pilotage. Les 16 tests exécutés, représentant 98 assertions, ont réussi. Ce résultat confirme les comportements couverts par ces scénarios, sans remplacer une expérimentation avec des usagers ni une observation de la charge de travail après mise en service.")
add_heading("4.7 Limites constatées et suggestions", 4)
add_paragraph("Le premier point à harmoniser concerne le délai. Le fonctionnement métier communiqué par l’ANBG fixe un objectif de 72 heures ouvrées. Dans la version du code vérifiée, le calcul de suivi et d’alerte est paramétré par défaut sur 24 heures ouvrées, tandis que l’accusé de réception adressé à l’usager annonce 72 heures. Avant une généralisation, une règle unique doit être validée par l’UCAS et le contrôle interne, puis appliquée aux messages, aux indicateurs et aux tests.")
add_paragraph("Le second point touche au circuit de réponse. Dans la pratique antérieure, l’agent répond au chef de service, le chef renvoie sa réponse à l’UCAS, et seule l’UCAS transmet la réponse finale à l’usager. L’application permet des actions de réponse à plusieurs rôles. Il convient donc de vérifier avec les responsables si la validation et l’envoi final doivent rester exclusivement confiés à l’UCAS, puis de configurer les droits et le parcours en conséquence.")
add_paragraph("La transition depuis Gmail et Excel demande également un accompagnement : définir comment les réclamations reçues par courriel seront intégrées, décider si l’historique Excel doit être repris, former les agents concernés et désigner une personne chargée de surveiller les échecs d’envoi des courriels. Après une période pilote, les indicateurs de volume, de délai et de répartition par service permettront d’évaluer le bénéfice réel de l’outil et d’ajuster les procédures.")
add_paragraph("Enfin, les demandes d’information apparaissent dans l’étude des besoins, alors que le formulaire public de la version vérifiée est centré sur les réclamations. Si ces demandes doivent entrer dans le périmètre de l’application, leur dépôt et leur traitement devront être ajoutés ; sinon, les documents fonctionnels devront être harmonisés avec le périmètre effectivement retenu.")
add_heading("Conclusion du chapitre IV", 4)
add_paragraph("L’application répond aux principaux besoins de traçabilité, de coordination et de reporting identifiés lors de l’étude. Les vérifications techniques sont encourageantes, mais la cohérence des règles de délai, la validation du circuit de réponse et l’accompagnement des utilisateurs conditionnent la qualité de son exploitation.")

add_heading("CONCLUSION GÉNÉRALE", 1, page_break=True)
add_paragraph("Ce stage à l’ANBG m’a permis de découvrir l’organisation de l’établissement et le travail de son service informatique. Les activités de déploiement de postes, d’imprimantes et d’onduleurs, la configuration d’équipements, l’intégration d’ordinateurs au domaine et la rédaction de procès-verbaux de réunions ont complété l’étude menée autour de la gestion des réclamations.")
add_paragraph("L’analyse du traitement par courriel et du reporting sur Excel a mis en évidence l’intérêt d’un espace commun pour enregistrer, affecter et suivre les dossiers. La solution conçue traduit ce besoin par des rôles distincts, des étapes de traitement tracées et des outils de pilotage. Sa modélisation et sa mise en œuvre constituent le résultat central du projet présenté dans ce rapport.")
add_paragraph("La poursuite du travail devra porter sur la validation des règles métier avec les services concernés, notamment le délai de 72 heures ouvrées et le rôle de l’UCAS dans l’envoi de la réponse finale. Une utilisation pilote, accompagnée d’une mesure des délais et de la charge de reporting, permettra ensuite d’apprécier les effets de l’application en situation réelle.")

add_heading("BIBLIOGRAPHIE ET DOCUMENTS DE RÉFÉRENCE", 1, page_break=True)
for reference in [
    "ANBG, Cahier des charges — Définition du suivi des réclamations, document interne communiqué pendant le stage.",
    "ANBG, Décret n° 0003/PR/MESRSTTENFC du 11 janvier 2021, document institutionnel fourni pour la présentation de l’organisme.",
    "IITG, Méthodologie de rédaction du rapport de stage, document pédagogique communiqué pour l’année académique 2025–2026.",
    "IITG, Gabarit de mise en page du rapport de stage 2026, version PDF fournie par l’étudiant.",
    "Projet Suivi-réclamation, code source de l’application et tests automatisés consultés le 20 septembre 2026.",
]:
    add_paragraph(reference)

add_heading("ANNEXE : VÉRIFICATION TECHNIQUE DE L’APPLICATION", 1, page_break=True)
add_paragraph("Synthèse des tests automatisés exécutés sur la version du projet consultée pour préparer ce rapport. Ces résultats décrivent une vérification technique ponctuelle et ne constituent pas une mesure d’usage en production.")
tests = doc.add_table(rows=1, cols=3)
tests.style = "Table Grid"
tests.rows[0].cells[0].text = "Groupe de tests"
tests.rows[0].cells[1].text = "Nombre"
tests.rows[0].cells[2].text = "Objet principal"
for row in [
    ("Parcours des réclamations", "10", "Affectation, traitement, réponse et courriels"),
    ("Règles de délai ouvré", "5", "Calcul des échéances et alertes"),
    ("Export de pilotage", "1", "Génération du document Word"),
    ("Total", "16", "98 assertions réussies"),
]:
    cells = tests.add_row().cells
    for cell, value in zip(cells, row):
        cell.text = value
for i, row in enumerate(tests.rows):
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.line_spacing = 1.1
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9.5)
                if i == 0:
                    run.bold = True

add_heading("TABLE DES MATIÈRES", 1, page_break=True)
contents = add_paragraph()
add_field(contents, 'TOC \\o "1-5" \\h \\z \\u', "Mettre à jour la table des matières dans Word.")

# Remove only empty spacer paragraphs from the inherited document, never
# paragraphs carrying figures, shapes, fields or page/section breaks.
for p in original:
    if p.text.strip():
        continue
    xml = p._p.xml
    if any(tag in xml for tag in ("w:drawing", "w:pict", "w:br", "w:sectPr", "w:fldChar")):
        continue
    parent = p._p.getparent()
    if parent is not None:
        parent.remove(p._p)

doc.core_properties.title = "Rapport de stage — Gestion des réclamations à l’ANBG"
doc.core_properties.subject = "Rapport de stage IITG"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
doc.save(OUT)
print(OUT)
